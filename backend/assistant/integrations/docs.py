"""Portable DOCX and PDF exporters; findings retain verification labels."""
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

def generate_report(report, output_dir):
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    paths = {kind: folder / f"report.{kind}" for kind in ("docx", "pdf")}
    doc = Document()
    doc.core_properties.author = "Derek"
    doc.core_properties.subject = "Research report"
    doc.add_heading(report["title"], 0)
    blocks = [("Executive summary", report.get("executive_summary", ""))]
    from .project_ideas import project_blocks
    blocks.extend(project_blocks(report))
    for i, finding in enumerate(report.get("key_findings", []), 1):
        label = "Corroborated" if finding.get("verified") else "Unverified"
        text = f"{label}: {finding['text']}\n" + "\n".join(finding.get("source_urls", []))
        blocks.append((f"Finding {i}", text))
    blocks.append(("Next steps", "\n".join(report.get("next_steps", []))))
    blocks.append(("Sources", "\n\n".join(f"{s.get('title', s['url'])}\n{s['url']}\nRetrievable: {'yes' if s.get('reachable') else 'not confirmed'}" for s in report.get("sources", []))))
    blocks.append(("Verification note", "Automated evidence checking is not a guarantee of truth. Findings without supporting evidence from two independent domains remain unverified."))
    styles = getSampleStyleSheet()
    styles["Heading2"].keepWithNext = True
    styles["Normal"].wordWrap = "CJK"
    story = [Paragraph(escape(report["title"]), styles["Title"]), Spacer(1, 0.2 * inch)]
    for heading, text in blocks:
        doc.add_heading(heading, level=1)
        for paragraph in str(text).split("\n"):
            doc.add_paragraph(paragraph)
        section = [Paragraph(escape(heading), styles["Heading2"])]
        for paragraph in str(text).split("\n"):
            section.append(Paragraph(escape(paragraph) or "&#160;", styles["Normal"]))
            section.append(Spacer(1, 5))
        if heading.startswith("Finding "):
            story.append(KeepTogether(section))
        else:
            story.extend(section)
    doc.save(paths["docx"])
    def footer(canvas, document):
        canvas.setFillColor(colors.HexColor("#52616b"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(inch * .7, inch * .45, "Derek • Research report")
        canvas.drawRightString(document.pagesize[0] - inch * .7, inch * .45, f"Page {document.page}")
    SimpleDocTemplate(str(paths["pdf"]), rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.7*inch, bottomMargin=.7*inch).build(story, onFirstPage=footer, onLaterPages=footer)
    return {key: str(value) for key, value in paths.items()}
